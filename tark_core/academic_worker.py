"""TARK Academic Worker: Autonomous Document & School/College Work Generator.

Generates beautifully formatted Word (.docx) notes, assignments, PowerPoint (.pptx)
presentations, practice test papers, and study schedules, and automatically opens
them in Microsoft Office / default desktop applications.
"""
from __future__ import annotations

from datetime import datetime
import json
import os
from pathlib import Path
import re
import subprocess
import sys
from typing import Any

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from pptx import Presentation
from pptx.util import Inches as PptxInches, Pt as PptxPt
from pptx.dml.color import RGBColor as PptxRGBColor
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from groq import Groq
from tark_core.config import (
    BASE_DIR,
    FALLBACK_REASONING_MODEL,
    GROQ_API_KEY,
    PRIMARY_REASONING_MODEL,
)

VAULT_DIR = BASE_DIR / "data" / "student_vault"
VAULT_DIR.mkdir(parents=True, exist_ok=True)


def _safe_filename(text: str) -> str:
    cleaned = re.sub(r"[^\w\s-]", "", text).strip()
    return re.sub(r"[-\s]+", "_", cleaned)[:50]


def _launch_file(filepath: Path) -> None:
    """Opens file in Microsoft Office or default Windows application."""
    try:
        if sys.platform == "win32":
            os.startfile(str(filepath))
        else:
            subprocess.Popen(["xdg-open", str(filepath)])
    except Exception as e:
        print(f"[AcademicWorker] Notice: Could not open {filepath}: {e}")


def _set_cell_background(cell, hex_color: str):
    """Sets background shading of a docx table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tc_pr.append(shd)


class AcademicWorker:
    """Specialist worker executing real academic document creation for students."""

    def __init__(self):
        self.groq_client = Groq(api_key=GROQ_API_KEY)

    def _query_llm(self, system_prompt: str, user_prompt: str) -> str:
        models = [PRIMARY_REASONING_MODEL, FALLBACK_REASONING_MODEL, "llama-3.1-8b-instant"]
        for model in models:
            try:
                resp = self.groq_client.chat.completions.create(
                    model=model,
                    messages=[
                        {"role": "system", "content": system_prompt},
                        {"role": "user", "content": user_prompt},
                    ],
                    temperature=0.3,
                    max_tokens=3000,
                )
                return resp.choices[0].message.content or ""
            except Exception as e:
                print(f"[AcademicWorker] Model {model} error: {e}")
                continue
        return "Unable to generate academic content at this time."

    # =========================================================================
    # 1. Word Revision Notes (.docx)
    # =========================================================================
    def create_word_notes(
        self,
        topic: str,
        subject: str = "General",
        grade: str = "10",
        board: str = "CBSE",
        student_name: str = "Shishir",
        open_after_create: bool = True,
    ) -> dict[str, Any]:
        """Generates comprehensive revision notes in Word and launches it."""
        system_prompt = (
            "You are an elite academic curriculum author and tutor. Produce structured revision notes. "
            "Return raw JSON with keys:\n"
            "- 'title': concise main topic\n"
            "- 'summary': 2-3 sentence overview\n"
            "- 'core_formulae': list of objects with {'name': str, 'formula': str, 'explanation': str}\n"
            "- 'sections': list of objects with {'heading': str, 'points': [str]}\n"
            "- 'exam_traps': list of common mistakes and examiner tips\n"
            "Output strictly valid JSON with no markdown formatting."
        )
        user_prompt = f"Topic: {topic}\nSubject: {subject}\nGrade: {grade}\nBoard: {board}\nStudent: {student_name}"

        raw = self._query_llm(system_prompt, user_prompt)
        try:
            clean = raw.strip()
            if clean.startswith("```json"):
                clean = clean[7:]
            if clean.endswith("```"):
                clean = clean[:-3]
            data = json.loads(clean.strip())
        except Exception:
            data = {
                "title": topic,
                "summary": f"Comprehensive revision guide for {topic} ({subject}).",
                "core_formulae": [{"name": "Key Principle", "formula": "See theory", "explanation": "Core relationship"}],
                "sections": [{"heading": "Key Concepts", "points": [raw[:500]]}],
                "exam_traps": ["Memorize standard definitions and SI units carefully."],
            }

        # Build Word Document
        doc = Document()

        # Page Title
        title_p = doc.add_paragraph()
        title_run = title_p.add_run(data.get("title", topic).upper())
        title_run.font.size = Pt(22)
        title_run.font.bold = True
        title_run.font.color.rgb = RGBColor(16, 44, 87) # Deep Royal Navy
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Subtitle Meta
        meta_p = doc.add_paragraph()
        meta_run = meta_p.add_run(f"Subject: {subject} | Class: {grade} ({board}) | Student: {student_name} | Date: {datetime.now().strftime('%d %B %Y')}")
        meta_run.font.size = Pt(10)
        meta_run.font.italic = True
        meta_run.font.color.rgb = RGBColor(100, 110, 120)
        meta_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph()

        # Summary Callout Table
        summary_table = doc.add_table(rows=1, cols=1)
        summary_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = summary_table.cell(0, 0)
        _set_cell_background(cell, "F0F4F8")
        cell_p = cell.paragraphs[0]
        c_bold = cell_p.add_run("📌 CHAPTER OVERVIEW: ")
        c_bold.font.bold = True
        c_bold.font.color.rgb = RGBColor(16, 44, 87)
        c_text = cell_p.add_run(data.get("summary", ""))
        c_text.font.size = Pt(10.5)
        doc.add_paragraph()

        # Core Formulas / Laws Table
        formulas = data.get("core_formulae", [])
        if formulas:
            f_head = doc.add_paragraph()
            f_run = f_head.add_run("⚡ CORE FORMULAS & LAWS")
            f_run.font.size = Pt(13)
            f_run.font.bold = True
            f_run.font.color.rgb = RGBColor(16, 44, 87)

            f_table = doc.add_table(rows=1, cols=3)
            f_table.alignment = WD_TABLE_ALIGNMENT.CENTER
            hdr_cells = f_table.rows[0].cells
            hdr_cells[0].text = "Concept / Law"
            hdr_cells[1].text = "Formula / Statement"
            hdr_cells[2].text = "Significance & SI Units"
            for c in hdr_cells:
                _set_cell_background(c, "1E3A8A")
                for p in c.paragraphs:
                    for r in p.runs:
                        r.font.bold = True
                        r.font.color.rgb = RGBColor(255, 255, 255)
                        r.font.size = Pt(10)

            for f in formulas:
                row_cells = f_table.add_row().cells
                row_cells[0].text = str(f.get("name", ""))
                row_cells[1].text = str(f.get("formula", ""))
                row_cells[2].text = str(f.get("explanation", ""))
                for c in row_cells:
                    _set_cell_background(c, "F8FAFC")
                    for p in c.paragraphs:
                        for r in p.runs:
                            r.font.size = Pt(9.5)
            doc.add_paragraph()

        # Sections & Detailed Breakdown
        for sec in data.get("sections", []):
            s_head = doc.add_paragraph()
            s_run = s_head.add_run(f"📖 {sec.get('heading', 'Concepts')}")
            s_run.font.size = Pt(12.5)
            s_run.font.bold = True
            s_run.font.color.rgb = RGBColor(30, 58, 138)

            for pt in sec.get("points", []):
                p = doc.add_paragraph(style="List Bullet")
                p_run = p.add_run(pt)
                p_run.font.size = Pt(10.5)

        # Exam Traps & Tips
        traps = data.get("exam_traps", [])
        if traps:
            doc.add_paragraph()
            t_head = doc.add_paragraph()
            t_run = t_head.add_run("⚠️ EXAM PITFALLS & PREVIOUS YEAR TRAPS")
            t_run.font.size = Pt(12.5)
            t_run.font.bold = True
            t_run.font.color.rgb = RGBColor(185, 28, 28) # Red

            t_table = doc.add_table(rows=1, cols=1)
            t_cell = t_table.cell(0, 0)
            _set_cell_background(t_cell, "FEF2F2")
            t_p = t_cell.paragraphs[0]
            for tr in traps:
                p_trap = t_cell.add_paragraph(style="List Bullet")
                r_trap = p_trap.add_run(tr)
                r_trap.font.size = Pt(10)
                r_trap.font.color.rgb = RGBColor(127, 29, 29)

        filename = f"{_safe_filename(topic)}_Notes.docx"
        out_path = VAULT_DIR / filename
        doc.save(str(out_path))

        if open_after_create:
            _launch_file(out_path)

        return {
            "status": "success",
            "file_path": str(out_path),
            "file_name": filename,
            "title": data.get("title", topic),
            "summary": f"Created revision notes for {topic} in Word and launched application.",
        }

    # =========================================================================
    # 2. School / College Assignment & Lab Report (.docx)
    # =========================================================================
    def create_assignment_document(
        self,
        title: str,
        subject: str = "Science",
        instructions: str = "",
        student_name: str = "Shishir",
        open_after_create: bool = True,
    ) -> dict[str, Any]:
        """Drafts an academic assignment or lab report in Word and opens it."""
        system_prompt = (
            "You are an academic scholar drafting a full formal school or college assignment/lab report. "
            "Return raw JSON with keys:\n"
            "- 'title': full assignment title\n"
            "- 'objective': specific aim or purpose of the assignment\n"
            "- 'theory': 2-3 paragraphs of theoretical background\n"
            "- 'procedure_or_method': list of step-by-step methods\n"
            "- 'observations': list of dicts with {'parameter': str, 'value_or_reading': str, 'unit': str}\n"
            "- 'conclusion': analytical conclusion and findings\n"
            "- 'questions': list of 2-3 discussion/viva questions with answers\n"
            "Output strictly valid JSON with no markdown formatting."
        )
        user_prompt = f"Assignment Topic: {title}\nSubject: {subject}\nSpecific Instructions: {instructions}\nStudent: {student_name}"

        raw = self._query_llm(system_prompt, user_prompt)
        try:
            clean = raw.strip()
            if clean.startswith("```json"):
                clean = clean[7:]
            if clean.endswith("```"):
                clean = clean[:-3]
            data = json.loads(clean.strip())
        except Exception:
            data = {
                "title": title,
                "objective": f"To study and analyze {title}.",
                "theory": raw[:600],
                "procedure_or_method": ["Step 1: Setup the apparatus.", "Step 2: Take measurements.", "Step 3: Analyze data."],
                "observations": [{"parameter": "Reading 1", "value_or_reading": "Standard", "unit": "N/A"}],
                "conclusion": f"The experiment successfully demonstrated the principles of {title}.",
                "questions": [{"question": "What is the key takeaway?", "answer": "The experimental values conform to theory."}],
            }

        doc = Document()

        # Cover / Header
        h_p = doc.add_paragraph()
        h_run = h_p.add_run(data.get("title", title).upper())
        h_run.font.size = Pt(20)
        h_run.font.bold = True
        h_run.font.color.rgb = RGBColor(15, 23, 42)
        h_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        meta_p = doc.add_paragraph()
        meta_run = meta_p.add_run(f"Submitted by: {student_name} | Subject: {subject} | Date: {datetime.now().strftime('%d %B %Y')}")
        meta_run.font.size = Pt(10.5)
        meta_run.font.italic = True
        meta_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph()

        # 1. Objective
        obj_head = doc.add_paragraph()
        obj_run = obj_head.add_run("🎯 AIM / OBJECTIVE")
        obj_run.font.size = Pt(12)
        obj_run.font.bold = True
        obj_run.font.color.rgb = RGBColor(30, 41, 59)
        doc.add_paragraph(data.get("objective", ""))

        # 2. Theoretical Background
        th_head = doc.add_paragraph()
        th_run = th_head.add_run("📚 THEORETICAL BACKGROUND")
        th_run.font.size = Pt(12)
        th_run.font.bold = True
        th_run.font.color.rgb = RGBColor(30, 41, 59)
        doc.add_paragraph(data.get("theory", ""))

        # 3. Method / Procedure
        proc_head = doc.add_paragraph()
        proc_run = proc_head.add_run("⚙️ PROCEDURE / METHODOLOGY")
        proc_run.font.size = Pt(12)
        proc_run.font.bold = True
        proc_run.font.color.rgb = RGBColor(30, 41, 59)
        for idx, step in enumerate(data.get("procedure_or_method", []), 1):
            p = doc.add_paragraph()
            p.add_run(f"{idx}. ").bold = True
            p.add_run(step)

        # 4. Observation Table
        obs = data.get("observations", [])
        if obs:
            doc.add_paragraph()
            obs_head = doc.add_paragraph()
            obs_run = obs_head.add_run("📊 OBSERVATION & DATA TABLE")
            obs_run.font.size = Pt(12)
            obs_run.font.bold = True
            obs_run.font.color.rgb = RGBColor(30, 41, 59)

            table = doc.add_table(rows=1, cols=3)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            hdr = table.rows[0].cells
            hdr[0].text = "Parameter / Trail"
            hdr[1].text = "Observed Value"
            hdr[2].text = "Unit / Remark"
            for c in hdr:
                _set_cell_background(c, "334155")
                for p in c.paragraphs:
                    for r in p.runs:
                        r.font.bold = True
                        r.font.color.rgb = RGBColor(255, 255, 255)
            for row in obs:
                r_cells = table.add_row().cells
                r_cells[0].text = str(row.get("parameter", ""))
                r_cells[1].text = str(row.get("value_or_reading", ""))
                r_cells[2].text = str(row.get("unit", ""))
                for c in r_cells:
                    _set_cell_background(c, "F1F5F9")

        # 5. Conclusion
        doc.add_paragraph()
        c_head = doc.add_paragraph()
        c_run = c_head.add_run("✅ CONCLUSION & FINDINGS")
        c_run.font.size = Pt(12)
        c_run.font.bold = True
        c_run.font.color.rgb = RGBColor(16, 185, 129)
        doc.add_paragraph(data.get("conclusion", ""))

        # 6. Viva / Discussion Questions
        viva = data.get("questions", [])
        if viva:
            doc.add_paragraph()
            v_head = doc.add_paragraph()
            v_run = v_head.add_run("💡 DISCUSSION & VIVA VOCE QUESTIONS")
            v_run.font.size = Pt(12)
            v_run.font.bold = True
            for item in viva:
                p_q = doc.add_paragraph(style="List Bullet")
                p_q.add_run("Q: " + item.get("question", "")).bold = True
                p_a = doc.add_paragraph()
                p_a.add_run("   A: " + item.get("answer", "")).italic = True

        filename = f"{_safe_filename(title)}_Assignment.docx"
        out_path = VAULT_DIR / filename
        doc.save(str(out_path))

        if open_after_create:
            _launch_file(out_path)

        return {
            "status": "success",
            "file_path": str(out_path),
            "file_name": filename,
            "title": data.get("title", title),
            "summary": f"Drafted assignment '{title}' in Word and launched application.",
        }

    # =========================================================================
    # 3. PowerPoint Presentations (.pptx)
    # =========================================================================
    def create_presentation_slides(
        self,
        topic: str,
        num_slides: int = 5,
        subject: str = "Science",
        student_name: str = "Shishir",
        open_after_create: bool = True,
    ) -> dict[str, Any]:
        """Generates a PowerPoint presentation and launches PowerPoint."""
        system_prompt = (
            "You are a master presentation designer. Create clean, high-impact presentation slides. "
            "Return raw JSON with key 'slides': list of objects with {'title': str, 'bullets': [str]}\n"
            f"Generate exactly {num_slides} slides.\n"
            "Output strictly valid JSON with no markdown formatting."
        )
        user_prompt = f"Topic: {topic}\nSubject: {subject}\nPresenter: {student_name}"

        raw = self._query_llm(system_prompt, user_prompt)
        try:
            clean = raw.strip()
            if clean.startswith("```json"):
                clean = clean[7:]
            if clean.endswith("```"):
                clean = clean[:-3]
            data = json.loads(clean.strip())
            slides_data = data.get("slides", [])
        except Exception:
            slides_data = [
                {"title": topic, "bullets": [f"Presented by {student_name}", f"Subject: {subject}"]},
                {"title": "Overview", "bullets": ["Key concepts and background", "Core fundamentals"]},
                {"title": "Detailed Analysis", "bullets": ["Practical applications", "Real world impact"]},
                {"title": "Conclusion", "bullets": ["Summary of insights", "Q&A Session"]},
            ]

        prs = Presentation()
        # Title slide
        title_slide_layout = prs.slide_layouts[0]
        s0 = prs.slides.add_slide(title_slide_layout)
        s0.shapes.title.text = topic.title()
        s0.placeholders[1].text = f"Subject: {subject}\nPresented by: {student_name}\nTARK Jarvis Academic Intelligence"

        # Content slides
        bullet_slide_layout = prs.slide_layouts[1]
        for s_info in slides_data[1:]:
            slide = prs.slides.add_slide(bullet_slide_layout)
            slide.shapes.title.text = s_info.get("title", topic)
            tf = slide.placeholders[1].text_frame
            tf.word_wrap = True
            bullets = s_info.get("bullets", [])
            for idx, b in enumerate(bullets):
                if idx == 0:
                    tf.text = b
                else:
                    p = tf.add_paragraph()
                    p.text = b
                    p.level = 0

        filename = f"{_safe_filename(topic)}_Presentation.pptx"
        out_path = VAULT_DIR / filename
        prs.save(str(out_path))

        if open_after_create:
            _launch_file(out_path)

        return {
            "status": "success",
            "file_path": str(out_path),
            "file_name": filename,
            "title": topic,
            "slides_count": len(prs.slides),
            "summary": f"Generated {len(prs.slides)}-slide presentation for '{topic}' in PowerPoint and opened it.",
        }

    # =========================================================================
    # 4. Practice Mock Test Paper (.docx)
    # =========================================================================
    def create_practice_test_paper(
        self,
        topic: str,
        subject: str = "Science",
        grade: str = "10",
        marks: int = 25,
        open_after_create: bool = True,
    ) -> dict[str, Any]:
        """Generates a complete mock exam paper with solutions in Word."""
        system_prompt = (
            f"You are a senior exam paper setter for CBSE/ICSE Grade {grade}. "
            f"Set a balanced {marks}-mark practice test on the specified topic. "
            "Return raw JSON with keys:\n"
            "- 'title': Test title\n"
            "- 'time_allowed': e.g. '45 Minutes'\n"
            "- 'general_instructions': [str]\n"
            "- 'sections': list of {'section_name': str, 'questions': [{'q_no': int, 'question': str, 'marks': int, 'options': [str] | None}]}\n"
            "- 'solutions': list of {'q_no': int, 'answer': str, 'marking_breakdown': str}\n"
            "Output strictly valid JSON with no markdown formatting."
        )
        user_prompt = f"Topic: {topic}\nSubject: {subject}\nGrade: {grade}\nTotal Marks: {marks}"

        raw = self._query_llm(system_prompt, user_prompt)
        try:
            clean = raw.strip()
            if clean.startswith("```json"):
                clean = clean[7:]
            if clean.endswith("```"):
                clean = clean[:-3]
            data = json.loads(clean.strip())
        except Exception:
            data = {
                "title": f"Practice Test: {topic}",
                "time_allowed": "45 Minutes",
                "general_instructions": ["Attempt all questions.", "Write clear steps for calculations."],
                "sections": [{
                    "section_name": "Section A (Objective & Conceptual)",
                    "questions": [{"q_no": 1, "question": f"State the primary law governing {topic}.", "marks": 2, "options": None}],
                }],
                "solutions": [{"q_no": 1, "answer": "Standard definition as per syllabus.", "marking_breakdown": "1 mark statement + 1 mark units"}],
            }

        doc = Document()

        # Header
        h_p = doc.add_paragraph()
        h_run = h_p.add_run(f"PRACTICE EXAMINATION: {topic.upper()}")
        h_run.font.size = Pt(18)
        h_run.font.bold = True
        h_run.font.color.rgb = RGBColor(15, 23, 42)
        h_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        meta_p = doc.add_paragraph()
        meta_run = meta_p.add_run(f"Subject: {subject} | Class: {grade} | Maximum Marks: {marks} | Time Allowed: {data.get('time_allowed', '45 Mins')}")
        meta_run.font.size = Pt(10)
        meta_run.font.bold = True
        meta_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Instructions
        ins_p = doc.add_paragraph()
        ins_p.add_run("General Instructions:\n").bold = True
        for ins in data.get("general_instructions", []):
            ins_p.add_run(f"• {ins}\n")

        doc.add_paragraph("―" * 45)

        # Questions
        for sec in data.get("sections", []):
            s_p = doc.add_paragraph()
            s_run = s_p.add_run(sec.get("section_name", "Section").upper())
            s_run.font.size = Pt(12)
            s_run.font.bold = True
            s_run.font.color.rgb = RGBColor(30, 58, 138)

            for q in sec.get("questions", []):
                q_p = doc.add_paragraph()
                q_p.add_run(f"Q{q.get('q_no', '')}. ").bold = True
                q_p.add_run(f"{q.get('question', '')} ")
                q_m = q_p.add_run(f"[{q.get('marks', 1)} Marks]")
                q_m.font.bold = True
                q_m.font.color.rgb = RGBColor(100, 116, 139)

                opts = q.get("options")
                if opts:
                    for opt in opts:
                        opt_p = doc.add_paragraph(style="List Bullet")
                        opt_p.add_run(opt)

        # Solutions & Marking Scheme Page
        doc.add_page_break()
        sol_h = doc.add_paragraph()
        sol_run = sol_h.add_run("DETAILED ANSWER KEY & MARKING SCHEME")
        sol_run.font.size = Pt(16)
        sol_run.font.bold = True
        sol_run.font.color.rgb = RGBColor(16, 185, 129)
        sol_h.alignment = WD_ALIGN_PARAGRAPH.CENTER

        for sol in data.get("solutions", []):
            sp = doc.add_paragraph()
            sp.add_run(f"Ans {sol.get('q_no', '')}: ").bold = True
            sp.add_run(sol.get("answer", ""))
            mb = sp.add_run(f"\n[Marking Scheme: {sol.get('marking_breakdown', '')}]")
            mb.font.italic = True
            mb.font.color.rgb = RGBColor(71, 85, 105)

        filename = f"{_safe_filename(topic)}_Practice_Test.docx"
        out_path = VAULT_DIR / filename
        doc.save(str(out_path))

        if open_after_create:
            _launch_file(out_path)

        return {
            "status": "success",
            "file_path": str(out_path),
            "file_name": filename,
            "title": f"Practice Test: {topic}",
            "marks": marks,
            "summary": f"Created {marks}-mark practice test for '{topic}' in Word with solutions.",
        }

    # =========================================================================
    # 5. Study Revision Timetable (.xlsx)
    # =========================================================================
    def create_study_timetable(
        self,
        subjects: list[str] | None = None,
        target_exam: str = "10th Board Exams",
        open_after_create: bool = True,
    ) -> dict[str, Any]:
        """Generates an hourly revision planner in Excel and opens it."""
        if not subjects:
            subjects = ["Physics Mechanics", "Calculus & Math", "Organic Chemistry", "Biology", "English Literature"]

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Study Revision Plan"

        # Title
        ws.merge_cells("A1:G1")
        title_cell = ws["A1"]
        title_cell.value = f"TARK MASTER STUDY SCHEDULE - TARGET: {target_exam.upper()}"
        title_cell.font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
        title_cell.fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 40

        # Subtitle
        ws.merge_cells("A2:G2")
        sub_cell = ws["A2"]
        sub_cell.value = f"Personalized Academic Timetable for Shishir | Generated on {datetime.now().strftime('%d %B %Y')}"
        sub_cell.font = Font(name="Calibri", size=11, italic=True, color="475569")
        sub_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[2].height = 20

        # Headers
        headers = ["Time Slot", "Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday / Sunday"]
        for col_idx, h in enumerate(headers, 1):
            c = ws.cell(row=3, column=col_idx)
            c.value = h
            c.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            c.fill = PatternFill(start_color="3B82F6", end_color="3B82F6", fill_type="solid")
            c.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[3].height = 26

        # Slots
        schedule_rows = [
            ("06:00 AM - 07:30 AM", subjects[0], subjects[1], subjects[0], subjects[2], subjects[0], "Full Length Mock Exam"),
            ("07:30 AM - 08:30 AM", "Breakfast & Prep", "Breakfast & Prep", "Breakfast & Prep", "Breakfast & Prep", "Breakfast & Prep", "Morning Walk & Relaxation"),
            ("04:30 PM - 06:00 PM", subjects[1], subjects[2], subjects[1], subjects[3] if len(subjects) > 3 else subjects[0], subjects[1], "Mock Paper Review & Errors"),
            ("06:30 PM - 08:00 PM", subjects[2], subjects[0], subjects[3] if len(subjects) > 3 else subjects[2], subjects[1], subjects[2], "Formula Sheet & Flashcards"),
            ("08:30 PM - 09:30 PM", "Daily Revision & Quiz", "Daily Revision & Quiz", "Daily Revision & Quiz", "Daily Revision & Quiz", "Daily Revision & Quiz", "Weekly Milestone Review"),
        ]

        thin_border = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1'),
        )

        for r_idx, row in enumerate(schedule_rows, 4):
            ws.row_dimensions[r_idx].height = 24
            for c_idx, val in enumerate(row, 1):
                cell = ws.cell(row=r_idx, column=c_idx)
                cell.value = val
                cell.font = Font(name="Calibri", size=10)
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = thin_border
                if c_idx == 1:
                    cell.fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
                    cell.font = Font(name="Calibri", size=10, bold=True)
                else:
                    cell.fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

        # Column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 18)

        filename = "Master_Study_Timetable.xlsx"
        out_path = VAULT_DIR / filename
        wb.save(str(out_path))

        if open_after_create:
            _launch_file(out_path)

        return {
            "status": "success",
            "file_path": str(out_path),
            "file_name": filename,
            "summary": "Generated weekly study timetable in Excel and opened application.",
        }


academic_worker = AcademicWorker()
